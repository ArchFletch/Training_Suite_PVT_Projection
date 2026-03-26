"""Format-agnostic dataset directory scanner.

Scans an arbitrary directory tree and extracts metadata from every file it can
read — column names, array shapes, row counts, value previews — without
assuming any particular dataset layout.  The output is a flat list of
``FileInfo`` records that downstream consumers (the LLM classifier, the wizard
UI) can inspect to decide what role each file plays.

Supported file formats:
- CSV / TSV (with header row detection)
- Touchstone (.s1p … .s99p)
- Pickle (.pkl) containing numpy arrays
- NumPy (.npy / .npz)
- HDF5 (.h5 / .hdf5)
- MATLAB (.mat)
- Excel (.xlsx / .xls)
- Images (.png / .jpg / .bmp / .tiff)
- GDS (.gds)
- Plain text with list-per-line (log.txt style)
"""

from __future__ import annotations

import ast
import os
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any


# ---------------------------------------------------------------------------
# Public data structures
# ---------------------------------------------------------------------------

@dataclass
class FileInfo:
    """Metadata extracted from a single file or coherent group of files."""

    # Relative path from the scanned root (e.g. "tabular/oct_coil_L.csv").
    rel_path: str
    # Absolute path on disk.
    abs_path: str
    # File extension, lower-cased (e.g. ".csv", ".s4p", ".pkl").
    extension: str
    # Number of files when this entry represents a directory of similar files.
    # 1 for a single file.
    file_count: int = 1
    # True when this entry represents a *group* of similarly-named files in a
    # directory rather than a single file.
    is_group: bool = False
    # Column names extracted from a header row (CSV, TSV, Excel, etc.).
    columns: list[str] = field(default_factory=list)
    # Number of data rows (excluding header).
    row_count: int | None = None
    # Shape of a numeric array (pkl, npy, npz, mat, hdf5).
    shape: list[int] | None = None
    # Numpy dtype string, when available.
    dtype: str | None = None
    # A few sample values from the first row / first slice.
    preview_values: list[Any] = field(default_factory=list)
    # Value-range per column: {col_name: {"min": ..., "max": ...}}.
    value_ranges: dict[str, dict[str, float]] = field(default_factory=dict)
    # Columns whose values are constant across all rows.
    constant_columns: list[str] = field(default_factory=list)
    # Extra key-value metadata (touchstone header fields, etc.).
    extra: dict[str, Any] = field(default_factory=dict)
    # Human-readable one-line summary (filled by the scanner).
    summary: str = ""


@dataclass
class ScanResult:
    """Complete scan output for one dataset directory."""

    root: str
    files: list[FileInfo] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def scan_directory(
    directory: str | Path,
    *,
    max_depth: int = 3,
    group_threshold: int = 5,
) -> ScanResult:
    """Scan *directory* and return metadata for every readable file / group.

    Parameters
    ----------
    directory:
        Path to the dataset root folder.
    max_depth:
        Maximum directory nesting depth to traverse.
    group_threshold:
        When a directory contains at least this many files with the same
        extension, they are reported as a single ``FileInfo`` group instead
        of individual entries.
    """
    directory = Path(directory).resolve()
    if not directory.is_dir():
        raise FileNotFoundError(f"Not a directory: {directory}")

    result = ScanResult(root=str(directory))
    _walk(directory, directory, result, max_depth, group_threshold, depth=0)
    return result


# ---------------------------------------------------------------------------
# Recursive directory walker
# ---------------------------------------------------------------------------

def _walk(
    current: Path,
    root: Path,
    result: ScanResult,
    max_depth: int,
    group_threshold: int,
    depth: int,
) -> None:
    """Recursively walk *current*, analyzing files and grouping where useful."""
    if depth > max_depth:
        return

    entries = sorted(current.iterdir())
    files = [e for e in entries if e.is_file() and not e.name.startswith(".")]
    dirs = [e for e in entries if e.is_dir() and not e.name.startswith(".")]

    # Group files by extension.
    ext_groups: dict[str, list[Path]] = {}
    for f in files:
        ext = f.suffix.lower()
        ext_groups.setdefault(ext, []).append(f)

    for ext, group_files in ext_groups.items():
        if len(group_files) >= group_threshold:
            # Treat as a group — analyze just the first file, report count.
            info = _analyze_file(group_files[0], root)
            info.is_group = True
            info.file_count = len(group_files)
            info.rel_path = str(current.relative_to(root))
            info.abs_path = str(current)
            info.summary = (
                f"{len(group_files)} {ext} files in {info.rel_path}/"
                + (f" — columns: {info.columns}" if info.columns else "")
                + (f" — shape per file: {info.shape}" if info.shape else "")
            )
            result.files.append(info)
        else:
            for f in group_files:
                info = _analyze_file(f, root)
                result.files.append(info)

    # Recurse into subdirectories.
    for d in dirs:
        _walk(d, root, result, max_depth, group_threshold, depth + 1)


# ---------------------------------------------------------------------------
# Per-file analyzers
# ---------------------------------------------------------------------------

def _analyze_file(path: Path, root: Path) -> FileInfo:
    """Dispatch to the right analyzer based on file extension."""
    ext = path.suffix.lower()
    rel = str(path.relative_to(root))
    info = FileInfo(rel_path=rel, abs_path=str(path), extension=ext)

    try:
        if ext in (".csv", ".tsv"):
            _analyze_csv(path, info)
        elif re.fullmatch(r"\.s\d+p", ext):
            _analyze_touchstone(path, info)
        elif ext == ".pkl":
            _analyze_pkl(path, info)
        elif ext == ".npy":
            _analyze_npy(path, info)
        elif ext == ".npz":
            _analyze_npz(path, info)
        elif ext in (".h5", ".hdf5"):
            _analyze_hdf5(path, info)
        elif ext == ".mat":
            _analyze_mat(path, info)
        elif ext in (".xlsx", ".xls"):
            _analyze_excel(path, info)
        elif ext in (".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"):
            _analyze_image(path, info)
        elif ext == ".gds":
            _analyze_gds(path, info)
        elif ext == ".txt":
            _analyze_txt(path, info)
        else:
            info.summary = f"{ext} file"
    except Exception as exc:
        info.summary = f"Error reading: {exc}"
        info.extra["error"] = str(exc)

    if not info.summary:
        info.summary = _build_summary(info)

    return info


# --- CSV / TSV ---

def _analyze_csv(path: Path, info: FileInfo) -> None:
    sep = "\t" if info.extension == ".tsv" else ","
    with path.open(encoding="utf-8", errors="replace") as fh:
        header_line = fh.readline().strip()
        first_data = fh.readline().strip()
        row_count = 1  # first_data line
        for _ in fh:
            row_count += 1

    parts = [p.strip() for p in header_line.split(sep)]
    is_header = len(parts) >= 2 and all(not _is_number(p) for p in parts)

    if is_header:
        info.columns = parts
        info.row_count = row_count
        if first_data:
            info.preview_values = [p.strip() for p in first_data.split(sep)]
        # Deep analysis: value ranges + constant columns (only for smallish files).
        if row_count <= 50_000:
            _csv_value_analysis(path, info, sep)
    else:
        # No header — treat first line as data too.
        info.row_count = row_count + 1
        info.extra["has_header"] = False
        info.preview_values = [p.strip() for p in header_line.split(sep)]

    info.summary = _build_summary(info)


def _csv_value_analysis(path: Path, info: FileInfo, sep: str) -> None:
    """Compute value ranges and detect constant columns for a CSV."""
    import numpy as np

    rows: list[list[float]] = []
    id_col_idx: int | None = None
    with path.open(encoding="utf-8", errors="replace") as fh:
        fh.readline()  # skip header
        for line in fh:
            vals = line.strip().split(sep)
            numeric_row: list[float] = []
            for j, v in enumerate(vals):
                v = v.strip()
                if _is_number(v):
                    numeric_row.append(float(v))
                else:
                    numeric_row.append(float("nan"))
                    if id_col_idx is None:
                        id_col_idx = j
            rows.append(numeric_row)

    if not rows:
        return

    arr = np.array(rows)
    for j, col in enumerate(info.columns):
        if j >= arr.shape[1]:
            break
        col_data = arr[:, j]
        valid = col_data[~np.isnan(col_data)]
        if len(valid) > 0:
            info.value_ranges[col] = {"min": float(valid.min()), "max": float(valid.max())}
            if len(valid) > 1 and valid.std() < 1e-8:
                info.constant_columns.append(col)

    if id_col_idx is not None and id_col_idx < len(info.columns):
        info.extra["id_column"] = info.columns[id_col_idx]


# --- Touchstone ---

def _analyze_touchstone(path: Path, info: FileInfo) -> None:
    port_match = re.fullmatch(r"\.s(\d+)p", info.extension)
    ports = int(port_match.group(1)) if port_match else None
    info.extra["port_count"] = ports

    with path.open(encoding="utf-8", errors="replace") as fh:
        for line in fh:
            line = line.strip()
            if not line or line.startswith("!"):
                continue
            if line.startswith("#"):
                parts = line[1:].split()
                if len(parts) >= 3:
                    info.extra["frequency_unit"] = parts[0]
                    info.extra["data_format"] = parts[2]
            elif line.lower().startswith("[number of ports]"):
                info.extra["port_count"] = int(line.split("]", 1)[1].strip())
            elif line.lower().startswith("[number of frequencies]"):
                info.extra["num_frequencies"] = int(line.split("]", 1)[1].strip())
                info.row_count = info.extra["num_frequencies"]
            elif line.lower().startswith("[end]"):
                break
            elif not line.startswith("["):
                try:
                    info.extra["first_frequency"] = float(line.split()[0])
                except (ValueError, IndexError):
                    pass
                break

    if ports:
        sparams = [f"S{i}{j}" for i in range(1, ports + 1) for j in range(i, ports + 1)]
        info.extra["sparams"] = sparams
        info.columns = sparams

    info.summary = _build_summary(info)


# --- Pickle ---

def _analyze_pkl(path: Path, info: FileInfo) -> None:
    import pickle
    with path.open("rb") as fh:
        data = pickle.load(fh)
    info.extra["python_type"] = type(data).__name__
    if hasattr(data, "shape"):
        info.shape = list(data.shape)
        info.dtype = str(data.dtype) if hasattr(data, "dtype") else None
    elif isinstance(data, (list, tuple)):
        info.extra["length"] = len(data)


# --- NumPy ---

def _analyze_npy(path: Path, info: FileInfo) -> None:
    import numpy as np
    data = np.load(path, mmap_mode="r")
    info.shape = list(data.shape)
    info.dtype = str(data.dtype)


def _analyze_npz(path: Path, info: FileInfo) -> None:
    import numpy as np
    with np.load(path, allow_pickle=False) as data:
        info.extra["arrays"] = {
            name: {"shape": list(data[name].shape), "dtype": str(data[name].dtype)}
            for name in data.files
        }
    info.summary = f"npz with arrays: {list(info.extra['arrays'].keys())}"


# --- HDF5 ---

def _analyze_hdf5(path: Path, info: FileInfo) -> None:
    try:
        import h5py
    except ImportError:
        info.summary = "HDF5 file (h5py not installed)"
        return
    with h5py.File(path, "r") as f:
        datasets = {}
        def _visit(name, obj):
            if isinstance(obj, h5py.Dataset):
                datasets[name] = {"shape": list(obj.shape), "dtype": str(obj.dtype)}
        f.visititems(_visit)
        info.extra["datasets"] = datasets
    info.summary = f"HDF5 with datasets: {list(info.extra.get('datasets', {}).keys())}"


# --- MATLAB ---

def _analyze_mat(path: Path, info: FileInfo) -> None:
    try:
        from scipy.io import loadmat
    except ImportError:
        info.summary = ".mat file (scipy not installed)"
        return
    mat = loadmat(path, squeeze_me=True)
    variables = {}
    for key, val in mat.items():
        if key.startswith("_"):
            continue
        if hasattr(val, "shape"):
            variables[key] = {"shape": list(val.shape), "dtype": str(val.dtype)}
        else:
            variables[key] = {"type": type(val).__name__}
    info.extra["variables"] = variables
    info.summary = f".mat with variables: {list(variables.keys())}"


# --- Excel ---

def _analyze_excel(path: Path, info: FileInfo) -> None:
    try:
        import openpyxl
    except ImportError:
        info.summary = "Excel file (openpyxl not installed)"
        return
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = wb.active
    rows = list(sheet.iter_rows(max_row=2, values_only=True))
    wb.close()
    if rows:
        info.columns = [str(c) if c is not None else "" for c in rows[0]]
        info.row_count = sheet.max_row - 1 if sheet.max_row else 0
        if len(rows) > 1:
            info.preview_values = [str(c) if c is not None else "" for c in rows[1]]


# --- Images ---

def _analyze_image(path: Path, info: FileInfo) -> None:
    try:
        from PIL import Image
        with Image.open(path) as img:
            info.shape = [img.height, img.width]
            info.extra["mode"] = img.mode
    except ImportError:
        # Fall back to just reporting the extension.
        pass
    info.summary = f"Image {info.extension}"
    if info.shape:
        info.summary += f" {info.shape[1]}x{info.shape[0]}"


# --- GDS ---

def _analyze_gds(path: Path, info: FileInfo) -> None:
    info.summary = "GDS layout file"
    info.extra["file_size_kb"] = round(path.stat().st_size / 1024, 1)


# --- Plain text (log.txt style) ---

def _analyze_txt(path: Path, info: FileInfo) -> None:
    header_comment = None
    first_data_line = None
    data_line_count = 0
    with path.open(encoding="utf-8", errors="replace") as fh:
        for raw in fh:
            line = raw.strip()
            if not line:
                continue
            if line.startswith("#"):
                header_comment = line
                continue
            if first_data_line is None:
                first_data_line = line
            data_line_count += 1

    info.row_count = data_line_count

    # Try to parse as Python list-per-line (log.txt format).
    if first_data_line:
        try:
            row = ast.literal_eval(first_data_line)
            if isinstance(row, (list, tuple)):
                info.extra["columns_per_row"] = len(row)
                info.preview_values = list(row)
                if header_comment:
                    # Header comment may contain column names.
                    info.extra["header_comment"] = header_comment
                    # Try to extract names: "# col1, col2, col3" or "# col1 col2 col3"
                    stripped = header_comment.lstrip("#").strip()
                    if "," in stripped:
                        info.columns = [c.strip() for c in stripped.split(",")]
                    else:
                        info.columns = stripped.split()
        except (ValueError, SyntaxError):
            info.extra["first_line"] = first_data_line[:200]


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _is_number(s: str) -> bool:
    try:
        float(s)
        return True
    except (ValueError, TypeError):
        return False


def _build_summary(info: FileInfo) -> str:
    """Build a human-readable one-line summary from a FileInfo."""
    parts: list[str] = []
    if info.is_group:
        parts.append(f"{info.file_count} {info.extension} files in {info.rel_path}/")
    else:
        parts.append(info.rel_path)

    if info.columns:
        cols_str = ", ".join(info.columns[:6])
        if len(info.columns) > 6:
            cols_str += f", ... ({len(info.columns)} total)"
        parts.append(f"columns: [{cols_str}]")

    if info.row_count is not None:
        parts.append(f"{info.row_count} rows")

    if info.shape:
        parts.append(f"shape: {info.shape}")

    if info.dtype:
        parts.append(f"dtype: {info.dtype}")

    return " — ".join(parts)


def format_scan_for_llm(scan: ScanResult) -> str:
    """Format a ScanResult into a concise text block for LLM classification."""
    lines = [f"Dataset directory: {scan.root}", ""]
    for i, f in enumerate(scan.files):
        label = f"[{i}]"
        lines.append(f"{label} {f.summary}")
        if f.columns and not f.is_group:
            lines.append(f"     columns: {f.columns}")
        if f.value_ranges:
            for col, vr in list(f.value_ranges.items())[:6]:
                lines.append(f"     {col}: {vr['min']:.4g} .. {vr['max']:.4g}")
        if f.constant_columns:
            lines.append(f"     constant: {f.constant_columns}")
        if f.extra.get("sparams"):
            lines.append(f"     S-parameters: {f.extra['sparams']}")
        if f.extra.get("arrays"):
            for name, meta in f.extra["arrays"].items():
                lines.append(f"     {name}: shape={meta['shape']}, dtype={meta['dtype']}")
        if f.extra.get("datasets"):
            for name, meta in f.extra["datasets"].items():
                lines.append(f"     {name}: shape={meta['shape']}, dtype={meta['dtype']}")
        if f.extra.get("variables"):
            for name, meta in f.extra["variables"].items():
                lines.append(f"     {name}: {meta}")
    return "\n".join(lines)
